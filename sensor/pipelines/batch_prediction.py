from sensor.exception import SensorException
from sensor.logger import logging
from sensor.predictors import ModelResolver
import pandas as pd
from sensor.utils import load_object
import os, sys
from datetime import datetime

PREDICTION_DIR = "prediction"

import numpy as np


def start_batch_prediction(input_file_path):
    try:
        os.makedirs(PREDICTION_DIR, exist_ok=True)
        logging.info(f"Creating model resolver object")
        model_resolver = ModelResolver(model_registry="saved_models")
        logging.info(f"Reading file :{input_file_path}")
        df = pd.read_csv(input_file_path)
        df.replace({"na": np.nan}, inplace=True)
        # validation

        logging.info(f"Loading transformer to transform dataset")
        transformer = load_object(
            file_path=model_resolver.get_latest_transformer_path()
        )

        input_feature_names = list(transformer.feature_names_in_)
        input_arr = transformer.transform(df[input_feature_names])

        logging.info(f"Loading model to make prediction")
        model = load_object(file_path=model_resolver.get_latest_model_path())
        prediction = model.predict(input_arr)

        logging.info(f"Target encoder to convert predicted column into categorical")
        target_encoder = load_object(
            file_path=model_resolver.get_latest_target_encoder_path()
        )

        cat_prediction = target_encoder.inverse_transform(prediction)

        df["prediction"] = prediction
        df["cat_pred"] = cat_prediction

        prediction_file_name = os.path.basename(input_file_path).replace(
            ".csv", f"{datetime.now().strftime('%m%d%Y__%H%M%S')}.csv"
        )
        prediction_file_path = os.path.join(PREDICTION_DIR, prediction_file_name)
        df.to_csv(prediction_file_path, index=False, header=True)

        # Feature-wise summary report (as DataFrame)
        summary = []
        for col in df.columns:
            if col not in ["prediction", "cat_pred"]:
                for pred_class in df["cat_pred"].unique():
                    stats = df[df["cat_pred"] == pred_class][col].describe()
                    summary.append({
                        "feature": col,
                        "predicted_class": pred_class,
                        "count": stats["count"],
                        "mean": stats["mean"] if "mean" in stats else None,
                        "std": stats["std"] if "std" in stats else None,
                        "min": stats["min"] if "min" in stats else None,
                        "max": stats["max"] if "max" in stats else None,
                    })
        summary_df = pd.DataFrame(summary)

        # --- Plots ---
        import matplotlib.pyplot as plt
        # Prediction distribution plot
        plt.figure(figsize=(6,4))
        df["cat_pred"].value_counts().plot(kind="bar")
        plt.title("Prediction Distribution")
        plt.xlabel("Predicted Class")
        plt.ylabel("Count")
        pred_dist_plot = os.path.join(PREDICTION_DIR, "prediction_distribution.png")
        plt.tight_layout()
        plt.savefig(pred_dist_plot)
        plt.close()
        # Feature distribution by class (for first 2 features)
        feature_dist_plots = []
        for col in df.columns:
            if col not in ["prediction", "cat_pred"] and pd.api.types.is_numeric_dtype(df[col]):
                plt.figure(figsize=(6,4))
                for cls in df["cat_pred"].unique():
                    df[df["cat_pred"]==cls][col].plot(kind="kde", label=str(cls))
                plt.title(f"Feature Distribution: {col}")
                plt.xlabel(col)
                plt.ylabel("Density")
                plt.legend()
                plot_path = os.path.join(PREDICTION_DIR, f"feature_dist_{col}.png")
                plt.tight_layout()
                plt.savefig(plot_path)
                plt.close()
                feature_dist_plots.append(plot_path)
                if len(feature_dist_plots) >= 2:
                    break

        # --- PDF Report ---
        from fpdf import FPDF
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", "B", 16)
        pdf.cell(0, 10, "Batch Prediction Report", ln=True, align="C")
        pdf.set_font("Arial", size=12)
        pdf.cell(0, 10, f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}", ln=True)
        pdf.ln(5)
        pdf.cell(0, 10, "Prediction Distribution:", ln=True)
        pdf.image(pred_dist_plot, w=100)
        pdf.ln(5)
        for plot in feature_dist_plots:
            pdf.cell(0, 10, f"Feature Distribution: {os.path.basename(plot).replace('feature_dist_','').replace('.png','')}", ln=True)
            pdf.image(plot, w=100)
            pdf.ln(3)
        pdf.ln(5)
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, "First 10 Predictions:", ln=True)
        pdf.set_font("Arial", size=8)
        # Table of first 10 predictions
        for idx, row in df.head(10).iterrows():
            row_str = ", ".join([f"{col}: {row[col]}" for col in df.columns])
            pdf.multi_cell(0, 6, row_str)
        pdf.ln(5)
        pdf.set_font("Arial", "B", 12)
        pdf.cell(0, 10, "Feature-wise Summary (first 10 rows):", ln=True)
        pdf.set_font("Arial", size=8)
        for idx, row in summary_df.head(10).iterrows():
            row_str = ", ".join([f"{col}: {row[col]}" for col in summary_df.columns])
            pdf.multi_cell(0, 6, row_str)
        pdf_file_path = os.path.join(PREDICTION_DIR, "prediction_report.pdf")
        pdf.output(pdf_file_path)

        # --- DOCX Report ---
        from docx import Document
        from docx.shared import Inches
        doc = Document()
        doc.add_heading("Batch Prediction Report", 0)
        doc.add_paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        doc.add_heading("Prediction Distribution", level=1)
        doc.add_picture(pred_dist_plot, width=Inches(4))
        for plot in feature_dist_plots:
            doc.add_heading(f"Feature Distribution: {os.path.basename(plot).replace('feature_dist_','').replace('.png','')}", level=2)
            doc.add_picture(plot, width=Inches(4))
        doc.add_heading("First 10 Predictions", level=1)
        table = doc.add_table(rows=1, cols=len(df.columns))
        hdr_cells = table.rows[0].cells
        for i, col in enumerate(df.columns):
            hdr_cells[i].text = str(col)
        for idx, row in df.head(10).iterrows():
            row_cells = table.add_row().cells
            for i, col in enumerate(df.columns):
                row_cells[i].text = str(row[col])
        doc.add_heading("Feature-wise Summary (first 10 rows)", level=1)
        table2 = doc.add_table(rows=1, cols=len(summary_df.columns))
        hdr_cells2 = table2.rows[0].cells
        for i, col in enumerate(summary_df.columns):
            hdr_cells2[i].text = str(col)
        for idx, row in summary_df.head(10).iterrows():
            row_cells2 = table2.add_row().cells
            for i, col in enumerate(summary_df.columns):
                row_cells2[i].text = str(row[col])
        docx_file_path = os.path.join(PREDICTION_DIR, "prediction_report.docx")
        doc.save(docx_file_path)

        return prediction_file_path
    except Exception as e:
        raise SensorException(e, sys)
